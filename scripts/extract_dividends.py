"""
Extract each mREIT's current QUARTERLY dividend per share from the workbook's
"Dividend" column (confirmed against its own "Div Yield on Price" column,
which computes (Dividend * 4) / Price) and save to data/dividends.csv:
    ticker, quarterly_dividend

This is a single current value per ticker, not a history -- a dividend cut or
raise makes the OLD value wrong going forward, so update it in the app's
Rankings tab (or via update_dividends.py) whenever a mREIT announces a change.

Re-running this script is safe: it only ADDS tickers missing from
data/dividends.csv (e.g. a newly tracked mREIT) and never overwrites a
ticker already present there, so it can't clobber a manually-updated
dividend with the workbook's stale snapshot value.

Same locked-OneDrive-file workaround as extract_bv.py (copy to temp, then read).
"""
import shutil
import tempfile
from pathlib import Path

import openpyxl
import pandas as pd

from common import DIVIDENDS_CSV, SOURCE_XLSX

SHEET_NAME = "mREITs"
TICKER_COL = 2       # column B
DIVIDEND_COL = 50    # column AX: quarterly Dividend per share
FIRST_DATA_ROW = 3


def load_workbook_safely(path: Path):
    """Copy the (possibly locked) workbook to a temp file, then load it."""
    with tempfile.TemporaryDirectory() as tmp:
        tmp_path = Path(tmp) / path.name
        shutil.copy2(path, tmp_path)
        return openpyxl.load_workbook(tmp_path, data_only=True)


def extract_dividends() -> int:
    wb = load_workbook_safely(SOURCE_XLSX)
    ws = wb[SHEET_NAME]

    existing_tickers = set()
    if DIVIDENDS_CSV.exists():
        existing_tickers = set(pd.read_csv(DIVIDENDS_CSV)["ticker"])

    new_rows = []
    r = FIRST_DATA_ROW
    while True:
        ticker = ws.cell(row=r, column=TICKER_COL).value
        if ticker is None:
            break
        ticker = str(ticker).strip().upper()
        if ticker not in existing_tickers:
            dividend = ws.cell(row=r, column=DIVIDEND_COL).value
            if dividend is not None:
                new_rows.append({"ticker": ticker, "quarterly_dividend": round(float(dividend), 4)})
        r += 1

    if not new_rows:
        print(f"No new tickers to add -- {DIVIDENDS_CSV} already covers everyone in the workbook.")
        return 0

    combined = pd.concat([pd.read_csv(DIVIDENDS_CSV), pd.DataFrame(new_rows)]) if existing_tickers else pd.DataFrame(new_rows)
    combined.to_csv(DIVIDENDS_CSV, index=False)
    print(f"Added {len(new_rows)} new dividend(s) ({[r['ticker'] for r in new_rows]}) to {DIVIDENDS_CSV}")
    return len(new_rows)


def main():
    extract_dividends()


if __name__ == "__main__":
    main()
