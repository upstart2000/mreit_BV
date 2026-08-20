"""
Extract book-value-per-share history for each mREIT from "BV Historical Data.xlsx"
and save it as a tidy long-format CSV at data/bv_quarterly.csv:
    ticker, quarter, quarter_end, book_value

The source workbook lives on OneDrive and is frequently locked (open in Excel /
mid-sync), so we always copy it to a temp file before reading with openpyxl.

Run this again any time the source workbook is updated with new quarters.
"""
import shutil
import tempfile
from pathlib import Path

import openpyxl
import pandas as pd

from common import BV_QUARTERLY_CSV, QUARTER_ENDS, QUARTER_ORDER, SOURCE_XLSX, SPLITS, normalize_quarter_label

SHEET_NAME = "mREITs"
TICKER_COL = 2          # column B
FIRST_BV_COL = 4        # column D (Q2'21 book value)
NUM_QUARTERS = len(QUARTER_ENDS)
HEADER_ROW = 2           # row with quarter labels
FIRST_DATA_ROW = 3


def load_workbook_safely(path: Path):
    """Copy the (possibly locked) workbook to a temp file, then load it."""
    with tempfile.TemporaryDirectory() as tmp:
        tmp_path = Path(tmp) / path.name
        shutil.copy2(path, tmp_path)
        return openpyxl.load_workbook(tmp_path, data_only=True)


def apply_split_adjustments(rows):
    """Multiply pre-split book values up to `last_pre_split_quarter` by `ratio`,
    in place, so they're on the same post-split share-count basis as the (always
    split-restated) daily prices. See SPLITS in common.py for why."""
    for ticker, split in SPLITS.items():
        cutoff_idx = QUARTER_ORDER.index(split["last_pre_split_quarter"])
        n_adjusted = 0
        for row in rows:
            if row["ticker"] != ticker:
                continue
            if QUARTER_ORDER.index(row["quarter"]) <= cutoff_idx:
                row["book_value"] = round(row["book_value"] * split["ratio"], 4)
                n_adjusted += 1
        print(
            f"Split-adjusted {ticker}: x{split['ratio']} on book values "
            f"through {split['last_pre_split_quarter']} ({n_adjusted} quarters)"
        )


def main():
    wb = load_workbook_safely(SOURCE_XLSX)
    ws = wb[SHEET_NAME]

    # Sanity check: read + normalize the BV quarter headers, confirm they match common.py
    header_labels = []
    for c in range(FIRST_BV_COL, FIRST_BV_COL + NUM_QUARTERS):
        raw = ws.cell(row=HEADER_ROW, column=c).value
        header_labels.append(normalize_quarter_label(raw))
    expected = list(QUARTER_ENDS.keys())
    if header_labels != expected:
        raise ValueError(
            f"BV quarter headers in workbook don't match common.QUARTER_ENDS.\n"
            f"Workbook: {header_labels}\nExpected: {expected}\n"
            "Update scripts/common.py QUARTER_ENDS if the workbook added new quarters."
        )

    rows = []
    r = FIRST_DATA_ROW
    while True:
        ticker = ws.cell(row=r, column=TICKER_COL).value
        if ticker is None:
            break
        ticker = str(ticker).strip().upper()
        for i, quarter in enumerate(expected):
            bv = ws.cell(row=r, column=FIRST_BV_COL + i).value
            if bv is None:
                continue
            rows.append(
                {
                    "ticker": ticker,
                    "quarter": quarter,
                    "quarter_end": QUARTER_ENDS[quarter],
                    # Rounded to avoid float-repr jitter (e.g. 15.549999999999999 vs
                    # 15.55) on re-extraction if Excel recalculates the source formulas.
                    "book_value": round(float(bv), 4),
                }
            )
        r += 1

    apply_split_adjustments(rows)

    # Preserve any quarters entered later through the app's "Update Book Values"
    # section (or scripts/update_bv.py) that aren't part of the fixed workbook set --
    # otherwise re-running this script would silently wipe them out.
    if BV_QUARTERLY_CSV.exists():
        existing = pd.read_csv(BV_QUARTERLY_CSV)
        extra = existing[~existing["quarter"].isin(expected)]
        if not extra.empty:
            rows.extend(extra.to_dict("records"))
            print(f"Preserved {len(extra)} manually-entered row(s) for quarter(s) not in the workbook: "
                  f"{sorted(extra['quarter'].unique())}")

    BV_QUARTERLY_CSV.parent.mkdir(exist_ok=True)
    with open(BV_QUARTERLY_CSV, "w", newline="", encoding="utf-8") as f:
        f.write("ticker,quarter,quarter_end,book_value\n")
        for row in rows:
            f.write(f"{row['ticker']},{row['quarter']},{row['quarter_end']},{row['book_value']}\n")

    tickers = sorted({row["ticker"] for row in rows})
    print(f"Extracted {len(rows)} (ticker, quarter) book-value rows for {len(tickers)} tickers:")
    print(", ".join(tickers))
    print(f"Saved to {BV_QUARTERLY_CSV}")


if __name__ == "__main__":
    main()
