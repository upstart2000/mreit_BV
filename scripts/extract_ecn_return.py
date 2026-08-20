"""
Extract quarterly economic return per mREIT from "BV Historical Data.xlsx"
and save it as a tidy long-format CSV at data/ecn_return_quarterly.csv:
    ticker, quarter, quarter_end, ecn_return

Trailing 1/2/3/4-year returns are NOT extracted as a static snapshot -- they're
computed on the fly (see common.trailing_return()) from this quarterly series,
so they automatically use whatever quarter is latest for each ticker and stay
current as new quarters are added, instead of going stale like the workbook's
own AT:AW columns would.

common.trailing_return() intentionally does NOT reproduce the workbook's own
2/3/4-Year Ecn Return columns exactly -- see its docstring for why (the
workbook's own multi-year window excludes the latest quarter; ours includes
it, at the user's request). Only the 1-Year convention is common to both, so
that's what gets sanity-checked below.

Same locked-OneDrive-file workaround as extract_bv.py (copy to temp, then read).
Run this again any time the source workbook is updated.
"""
import shutil
import tempfile
from pathlib import Path

import openpyxl

from common import ECN_RETURN_QUARTERLY_CSV, QUARTER_ENDS, SOURCE_XLSX, normalize_quarter_label, trailing_return

SHEET_NAME = "mREITs"
TICKER_COL = 2           # column B
FIRST_ECN_COL = 25       # column Y (Q2'21 economic return)
NUM_QUARTERS = len(QUARTER_ENDS)
ONE_YEAR_COL = 46        # column AT: workbook's own 1-Year Ecn Return, for a sanity check
LATEST_QUARTER = list(QUARTER_ENDS.keys())[-1]
HEADER_ROW = 2
FIRST_DATA_ROW = 3
TOLERANCE = 1e-4


def load_workbook_safely(path: Path):
    """Copy the (possibly locked) workbook to a temp file, then load it."""
    with tempfile.TemporaryDirectory() as tmp:
        tmp_path = Path(tmp) / path.name
        shutil.copy2(path, tmp_path)
        return openpyxl.load_workbook(tmp_path, data_only=True)


def extract_ecn_return() -> int:
    wb = load_workbook_safely(SOURCE_XLSX)
    ws = wb[SHEET_NAME]

    header_labels = []
    for c in range(FIRST_ECN_COL, FIRST_ECN_COL + NUM_QUARTERS):
        raw = ws.cell(row=HEADER_ROW, column=c).value
        header_labels.append(normalize_quarter_label(raw))
    expected = list(QUARTER_ENDS.keys())
    if header_labels != expected:
        raise ValueError(
            f"Ecn Retn quarter headers in workbook don't match common.QUARTER_ENDS.\n"
            f"Workbook: {header_labels}\nExpected: {expected}\n"
            "Update scripts/common.py QUARTER_ENDS if the workbook added new quarters."
        )

    rows = []
    workbook_1y = {}  # ticker -> workbook's own 1-Year Ecn Return, for the validation check
    r = FIRST_DATA_ROW
    while True:
        ticker = ws.cell(row=r, column=TICKER_COL).value
        if ticker is None:
            break
        ticker = str(ticker).strip().upper()
        for i, quarter in enumerate(expected):
            ret = ws.cell(row=r, column=FIRST_ECN_COL + i).value
            if ret is None:
                continue
            rows.append(
                {
                    "ticker": ticker,
                    "quarter": quarter,
                    "quarter_end": QUARTER_ENDS[quarter],
                    # Rounded to avoid float-repr jitter on re-extraction if Excel
                    # recalculates the source formulas (see extract_bv.py).
                    "ecn_return": round(float(ret), 6),
                }
            )
        workbook_1y[ticker] = ws.cell(row=r, column=ONE_YEAR_COL).value
        r += 1

    ECN_RETURN_QUARTERLY_CSV.parent.mkdir(exist_ok=True)
    with open(ECN_RETURN_QUARTERLY_CSV, "w", newline="", encoding="utf-8") as f:
        f.write("ticker,quarter,quarter_end,ecn_return\n")
        for row in rows:
            f.write(f"{row['ticker']},{row['quarter']},{row['quarter_end']},{row['ecn_return']}\n")

    tickers = sorted({row["ticker"] for row in rows})
    print(f"Extracted {len(rows)} (ticker, quarter) economic-return rows for {len(tickers)} tickers.")
    print(f"Saved to {ECN_RETURN_QUARTERLY_CSV}")

    validate_one_year(rows, workbook_1y)
    return len(rows)


def main():
    extract_ecn_return()


def validate_one_year(rows, workbook_1y):
    by_ticker = {}
    for row in rows:
        by_ticker.setdefault(row["ticker"], {})[row["quarter"]] = row["ecn_return"]

    known_quarters = {row["quarter"] for row in rows}
    mismatches = 0
    for ticker, returns in by_ticker.items():
        expected = workbook_1y.get(ticker)
        if expected is None:
            continue
        computed = trailing_return(returns, known_quarters, LATEST_QUARTER, 1)
        if computed is None or abs(computed - float(expected)) > TOLERANCE:
            mismatches += 1
            print(f"  MISMATCH {ticker} 1-Year: workbook={expected!r} computed={computed!r}")
    if mismatches:
        print(f"WARNING: {mismatches} 1-year mismatch(es) -- see above.")
    else:
        print("Validated: trailing_return(..., years=1) reproduces the workbook's 1-Year column.")


if __name__ == "__main__":
    main()
